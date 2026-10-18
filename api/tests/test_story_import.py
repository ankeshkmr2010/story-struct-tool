"""Format-agnostic import: adapters (Word, Scrivener) parse; one importer saves the story."""

import io
import zipfile

import pytest

from storytool.domain.story.scrivener_import import read_scrivener, rtf_to_text

BS = chr(92)
RQUOTE = chr(0x2019)


def rtf(*paragraphs: str) -> str:
    body = (BS + "par ").join(paragraphs)
    return "{" + BS + "rtf1" + BS + "ansi{" + BS + "fonttbl" + BS + "f0 Helvetica;}" + body + "}"


def binder_item(uuid: str, kind: str, title: str, children: str = "", compile_: bool = True):
    meta = "" if compile_ else "<MetaData><IncludeInCompile>No</IncludeInCompile></MetaData>"
    inner = f"<Children>{children}</Children>" if children else ""
    return (
        f'<BinderItem UUID="{uuid}" Type="{kind}"><Title>{title}</Title>{meta}{inner}</BinderItem>'
    )


def scriv_zip() -> bytes:
    draft = binder_item(
        "D",
        "DraftFolder",
        "Manuscript",
        binder_item(
            "C1",
            "Folder",
            "Chapter One",
            binder_item("S1", "Text", "Arrival")
            + binder_item("S2", "Text", "Cut scene", compile_=False)
            + binder_item("SUB", "Folder", "Part", binder_item("S3", "Text", "Nested")),
        )
        + binder_item("L1", "Text", "Interlude"),
    )
    research = binder_item("R", "ResearchFolder", "Research", binder_item("R1", "Text", "Notes"))
    project = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f"<ScrivenerProject><Binder>{draft}{research}</Binder></ScrivenerProject>"
    )
    files = {
        "Novel.scriv/Novel.scrivx": project,
        "Novel.scriv/Files/Data/S1/content.rtf": rtf("Maya arrived" + BS + "'92s late.", "Rain."),
        "Novel.scriv/Files/Data/S1/synopsis.txt": "Maya reaches the harbour.",
        "Novel.scriv/Files/Data/S1/notes.rtf": rtf("Check the tide tables."),
        "Novel.scriv/Files/Data/S2/content.rtf": rtf("Never compiled."),
        "Novel.scriv/Files/Data/S3/content.rtf": rtf("Deep inside."),
        "Novel.scriv/Files/Data/L1/content.rtf": rtf("A quiet pause."),
        "Novel.scriv/Files/Data/R1/content.rtf": rtf("Research only."),
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def test_rtf_to_text_keeps_prose_and_drops_formatting() -> None:
    source = (
        "{"
        + BS
        + "rtf1{"
        + BS
        + "colortbl;}Caf"
        + BS
        + "'e9 "
        + BS
        + "u8212"
        + BS
        + "'97 "
        + BS
        + "ldblquote hi"
        + BS
        + "rdblquote"
        + BS
        + "par{"
        + BS
        + "*"
        + BS
        + "annotation secret}End "
        + BS
        + "{x"
        + BS
        + "}}"
    )
    assert rtf_to_text(source) == "Café — “hi”\n\nEnd {x}"


def test_scrivener_draft_folders_become_chapters_and_scenes() -> None:
    story = read_scrivener(scriv_zip())
    assert [chapter.title for chapter in story.chapters] == ["Chapter One", "Interlude"]
    first = story.chapters[0]
    assert [scene.title for scene in first.scenes] == ["Arrival", "Nested"]
    assert first.scenes[0].prose == "Maya arrived" + RQUOTE + "s late.\n\nRain."
    assert first.scenes[0].summary == "Maya reaches the harbour."
    assert story.chapters[1].scenes[0].prose == "A quiet pause."
    assert "Research only." not in str(story) and "Never compiled." not in str(story)


@pytest.mark.parametrize(
    ("content", "message"),
    [(b"not a zip", "Choose a .zip"), (zipfile.ZipFile.__name__.encode(), "Choose a .zip")],
)
def test_scrivener_rejects_non_projects(content: bytes, message: str) -> None:
    with pytest.raises(ValueError, match=message):
        read_scrivener(content)


async def test_import_endpoint_saves_any_adapter_output(client) -> None:
    response = await client.post(
        "/api/imports/scrivener?title=Imported%20novel",
        content=scriv_zip(),
        headers={"Content-Type": "application/zip"},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["chapter_count"] == 2 and body["scene_count"] == 3
    sid = body["story"]["id"]
    scenes = (await client.get(f"/api/stories/{sid}/scenes")).json()
    arrival = next(scene for scene in scenes if scene["title"] == "Arrival")
    assert arrival["summary"] == "Maya reaches the harbour." and arrival["status"] == "drafted"
    unknown = await client.post(
        "/api/imports/pages?title=x", content=b"x", headers={"Content-Type": "application/zip"}
    )
    assert unknown.status_code == 404


def docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("The Lantern Keeper", style="Title")
    doc.add_paragraph("For Mara.")
    doc.add_paragraph("Chapter 1")
    doc.add_paragraph("Maya lit the lamp.")
    doc.add_paragraph("* * *")
    doc.add_paragraph("Morning came.")
    doc.add_paragraph("Chapter 2")
    doc.add_paragraph().add_run("Maya leaves the island.").italic = True
    doc.add_paragraph("The boat rocked.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_inspection_recommends_and_plans_choose_the_split() -> None:
    from storytool.domain.story.docx_import import DocxPlan, inspect_docx, read_docx

    data = docx_bytes()
    info = inspect_docx(data)
    assert info["stats"]["chapter_lines"] == 2 and info["stats"]["separators"] == 1
    assert info["recommended"] == {
        "chapter_split": "chapter_lines",
        "scene_split": "separators",
        "skip": [[0, 1]],
    }
    assert info["outline"][4]["cues"] == ["separator"]
    story = read_docx(data, DocxPlan(**info["recommended"], italic_synopsis=True))
    assert [c.title for c in story.chapters] == ["Chapter 1", "Chapter 2"]
    assert [s.prose for s in story.chapters[0].scenes] == ["Maya lit the lamp.", "Morning came."]
    assert story.chapters[1].scenes[0].summary == "Maya leaves the island."
    explicit = read_docx(
        data,
        DocxPlan(
            chapter_split="explicit",
            chapter_starts=[3],
            scene_split="none",
            skip=[[0, 2]],
            titles={3: "Lamplight"},
            summaries={3: "Night falls."},
        ),
    )
    assert len(explicit.chapters) == 1 and explicit.chapters[0].title == "Lamplight"
    assert explicit.chapters[0].scenes[0].summary == "Night falls."
    assert "Chapter 2" in explicit.chapters[0].scenes[0].prose  # explicit starts keep prose


def test_scrivener_plans_choose_levels_items_notes_and_excluded() -> None:
    from storytool.domain.story.scrivener_import import (
        ScrivenerPlan,
        inspect_scrivener,
        read_scrivener,
    )

    data = scriv_zip()
    info = inspect_scrivener(data)
    assert info["recommended"] == {"chapter_level": 1}
    assert info["stats"]["notes"] == 1 and info["stats"]["excluded_from_compile"] == 1
    assert [o["value"] for o in info["options"]["chapter_level"]] == [1, 2]
    deeper = read_scrivener(data, ScrivenerPlan(chapter_level=2))
    assert [c.title for c in deeper.chapters] == ["Arrival", "Part", "Interlude"]
    items = read_scrivener(data, ScrivenerPlan(chapter_items=["SUB"], include_notes=True))
    assert [c.title for c in items.chapters] == ["Arrival", "Part"]
    assert [s.title for s in items.chapters[1].scenes] == ["Nested", "Interlude"]
    assert items.chapters[0].scenes[0].notes == "Check the tide tables."
    everything = read_scrivener(
        data, ScrivenerPlan(include_excluded=True, synopsis_as_summary=False)
    )
    titles = [s.title for c in everything.chapters for s in c.scenes]
    assert "Cut scene" in titles and everything.chapters[0].scenes[0].summary is None


async def test_upload_preview_and_import_with_a_plan(client) -> None:
    created = await client.post(
        "/api/imports/docx/uploads?filename=Lantern.docx",
        content=docx_bytes(),
        headers={"Content-Type": "application/octet-stream"},
    )
    assert created.status_code == 201, created.text
    upload = created.json()
    assert upload["inspection"]["recommended"]["chapter_split"] == "chapter_lines"
    plan = {"chapter_split": "chapter_lines", "scene_split": "separators", "skip": [[0, 1]]}
    preview = await client.post(f"/api/imports/uploads/{upload['id']}/preview", json={"plan": plan})
    assert preview.status_code == 200, preview.text
    assert preview.json()["chapter_count"] == 2 and preview.json()["scene_count"] == 3
    bad = await client.post(
        f"/api/imports/uploads/{upload['id']}/preview", json={"plan": {"chapter_split": "x"}}
    )
    assert bad.status_code == 400 and "Invalid import plan" in bad.text
    done = await client.post(
        f"/api/imports/uploads/{upload['id']}/import", json={"title": "Lantern", "plan": plan}
    )
    assert done.status_code == 201, done.text
    status = (await client.get(f"/api/imports/uploads/{upload['id']}")).json()
    assert status["story_id"] == done.json()["story"]["id"] and status["inspection"] is None
    again = await client.post(
        f"/api/imports/uploads/{upload['id']}/import", json={"title": "Lantern", "plan": plan}
    )
    assert again.status_code == 400


async def test_mcp_tools_plan_and_import_an_upload(client) -> None:
    from tests.test_beta_ai import error_text, result
    from tests.test_oauth import approve, authorize, exchange, oauth_tool, register

    upload = (
        await client.post(
            "/api/imports/scrivener/uploads?filename=Novel.zip",
            content=scriv_zip(),
            headers={"Content-Type": "application/zip"},
        )
    ).json()

    async def token(scopes):
        registration = await register(client)
        request_id, verifier, _ = await authorize(client, registration)
        code = await approve(client, request_id, None, scopes)
        return (await exchange(client, registration, code, verifier)).json()["access_token"]

    reader = await token(["story:read", "library:read"])
    refused = error_text(
        await oauth_tool(client, reader, "get_manuscript_import", {"upload_id": upload["id"]})
    )
    assert "401" in refused or "403" in refused
    agent = await token(["story:read", "prose:read", "library:read", "story:create"])
    inspected = result(
        await oauth_tool(client, agent, "get_manuscript_import", {"upload_id": upload["id"]})
    )
    assert inspected["outline_total"] == len(inspected["inspection"]["outline"])
    assert "chapter_level" in inspected["inspection"]["guide"]
    plan = {"chapter_items": ["C1", "L1"], "titles": {"C1": "Harbour"}}
    previewed = result(
        await oauth_tool(
            client, agent, "preview_manuscript_import", {"upload_id": upload["id"], "plan": plan}
        )
    )
    assert [c["title"] for c in previewed["chapters"]] == ["Harbour", "Interlude"]
    imported = result(
        await oauth_tool(
            client,
            agent,
            "import_manuscript",
            {"upload_id": upload["id"], "title": "From my AI", "plan": plan},
        )
    )
    assert imported["chapter_count"] == 2
    status = (await client.get(f"/api/imports/uploads/{upload['id']}")).json()
    assert status["story_id"] == imported["story"]["id"]


def test_markdown_and_text_share_block_planning() -> None:
    from storytool.domain.story.block_import import inspect_blocks, read_blocks
    from storytool.domain.story.text_import import text_blocks

    source = (
        b"Title Page\n\n# Chapter One\n\nMaya lit the lamp.\n\n---\n\nMorning came.\n\n"
        b"Chapter Two\n===========\n\n*She leaves.*\n\nThe boat rocked.\f\nAfterword text."
    )
    blocks = text_blocks(source)
    assert [b.style for b in blocks][:2] == ["Normal", "Heading 1"]
    assert blocks[3].separator and blocks[5].style == "Heading 1" and blocks[6].italic
    assert blocks[-1].page_break_before
    info = inspect_blocks(blocks)
    assert info["recommended"]["chapter_split"] == "heading1"
    assert info["recommended"]["scene_split"] == "separators"
    story = read_blocks(blocks, None)
    assert [c.title for c in story.chapters] == ["Chapter One", "Chapter Two"]
    assert len(story.chapters[0].scenes) == 2


def epub_bytes(drm: bool = False) -> bytes:
    container = (
        '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>'
    )
    opf = (
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf"><manifest>'
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
        '<item id="c1" href="text/c1.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="c2" href="text/c2.xhtml" media-type="application/xhtml+xml"/>'
        '</manifest><spine><itemref idref="nav"/><itemref idref="c1"/><itemref idref="c2"/>'
        "</spine></package>"
    )
    page = '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>x</title></head><body>{}</body></html>'
    files = {
        "META-INF/container.xml": container,
        "OEBPS/content.opf": opf,
        "OEBPS/nav.xhtml": page.format("<nav><ol><li>Contents</li></ol></nav>"),
        "OEBPS/text/c1.xhtml": page.format(
            "<h1>Chapter One</h1><p>Maya&nbsp;lit the lamp.</p><p><em>A short opener.</em></p>"
        ),
        "OEBPS/text/c2.xhtml": page.format(
            "<h1>Chapter Two</h1><div><p>The boat rocked.</p></div>"
        ),
    }
    if drm:
        files["META-INF/encryption.xml"] = "<encryption><EncryptedData/></encryption>"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def test_epub_reads_spine_order_headings_and_refuses_drm() -> None:
    from storytool.domain.story.block_import import read_blocks
    from storytool.domain.story.epub_import import epub_blocks

    blocks = epub_blocks(epub_bytes())
    assert blocks[0].text == "Chapter One"
    assert "Contents" not in [b.text for b in blocks]
    assert blocks[2].italic and blocks[3].page_break_before
    story = read_blocks(blocks, None)
    assert [c.title for c in story.chapters] == ["Chapter One", "Chapter Two"]
    assert story.chapters[0].scenes[0].prose.startswith("Maya lit the lamp.")
    with pytest.raises(ValueError, match="DRM"):
        epub_blocks(epub_bytes(drm=True))


async def test_markdown_and_epub_uploads_preview_and_import(client) -> None:
    markdown = b"# One\n\nMaya lit the lamp.\n\n***\n\nMorning came.\n\n# Two\n\nThe boat rocked."
    for fmt, content, chapters, scenes in [
        ("text", markdown, 2, 3),
        ("epub", epub_bytes(), 2, 2),
    ]:
        created = await client.post(
            f"/api/imports/{fmt}/uploads?filename=book.{fmt}",
            content=content,
            headers={"Content-Type": "application/octet-stream"},
        )
        assert created.status_code == 201, (fmt, created.text)
        upload = created.json()
        plan = upload["inspection"]["recommended"]
        preview = await client.post(
            f"/api/imports/uploads/{upload['id']}/preview", json={"plan": plan}
        )
        assert preview.status_code == 200, (fmt, preview.text)
        assert (preview.json()["chapter_count"], preview.json()["scene_count"]) == (
            chapters,
            scenes,
        ), fmt
        done = await client.post(
            f"/api/imports/uploads/{upload['id']}/import", json={"title": fmt, "plan": plan}
        )
        assert done.status_code == 201, (fmt, done.text)
        assert done.json()["chapter_count"] == chapters
