"""Whole-story backup, restore-as-a-copy, and Word export. The words must be able to leave."""

import copy
import io
import re
import zipfile
from datetime import UTC, datetime
from typing import Any
from uuid import UUID, uuid5

from docx import Document
from docx.shared import Inches
from litestar import Controller, Request, Response, get, post
from litestar.exceptions import ClientException, NotFoundException
from sqlalchemy import select, text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession
from uuid_utils import uuid7

from storytool.domain.ai.commands import ENTITIES, LINKS
from storytool.domain.analysis import load_graph_by_id
from storytool.domain.auth.models import User
from storytool.domain.glossary import GlossaryEntry
from storytool.domain.narrative.prose import compile_manuscript
from storytool.domain.story.manuscript_formats import (
    Block,
    ImageFiles,
    blocks,
    manuscript_epub,
    manuscript_html,
    manuscript_text,
)
from storytool.domain.story.manuscript_pdf import manuscript_pdf
from storytool.domain.story.models import Story
from storytool.domain.story.schemas import StoryOut
from storytool.domain.story.timeline_controller import story_timeline
from storytool.domain.versioning.service import ADMIN, EXTRAS, checkpoint, full_state, restore_state

BACKUP_FORMAT = "storytool-backup"
MAX_BACKUP_BYTES = 200 * 1024 * 1024


def tables() -> dict[str, Any]:
    return {
        **{kind: model.__table__ for kind, (model, _, _) in ENTITIES.items()},
        **{kind: table for kind, (table, _, _, _, _) in LINKS.items()},
        **EXTRAS,
    }


def rehome(state: dict[str, Any], story_id: UUID, *, strict: bool = False,
           image_ids: dict[str, str] | None = None) -> dict[str, Any]:
    """Copy a full_state under fresh IDs. Strict mode rejects references outside the copy."""
    state = copy.deepcopy(state)
    identifiers = {row["id"]: str(uuid7()) for kind in ENTITIES for row in state[kind]}
    identifiers.update(image_ids or {r["id"]: str(uuid7()) for r in state.get("story_image", [])})
    identifiers[state["story"][0]["id"]] = str(story_id)
    for kind, table in tables().items():
        for row in state.get(kind, []):
            if "id" in row:
                row["id"] = identifiers[row["id"]] if kind in ENTITIES or kind == "story_image" else str(uuid7())
            for column in table.c:
                value = row.get(column.name)
                if not column.foreign_keys or value is None:
                    continue
                if value in identifiers:
                    row[column.name] = identifiers[value]
                elif strict:
                    raise ValueError(f"{kind}.{column.name} points outside this backup")
            if kind == "story_image":
                row["attachments"] = [{"entity": a["entity"], "id": identifiers[a["id"]]}
                                      for a in row.get("attachments", [])]
            if kind in {"scene", "scene_revision"} and row.get("content"):
                from storytool.domain.story.images import IMAGE_REF

                row["content"] = IMAGE_REF.sub(
                    lambda m: f"![{m[1]}](storytool-image:{identifiers.get(m[2], m[2])})", row["content"]
                )
    return state


def validated_state(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict) or payload.get("format") != BACKUP_FORMAT:
        raise ValueError("This is not a StoryTool backup file")
    if payload.get("version") != 1:
        raise ValueError("This backup was made by a newer StoryTool")
    state = payload.get("state")
    if not isinstance(state, dict) or state.get("format_version") != 1:
        raise ValueError("The backup's story data is missing or unreadable")
    known = tables()
    unknown = set(state) - set(known) - {"format_version"}
    if unknown:
        raise ValueError("Unknown sections: " + ", ".join(sorted(unknown)))
    clean: dict[str, Any] = {"format_version": 1}
    seen: set[str] = set()
    for kind, table in known.items():
        rows = state.get(kind, [])
        if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
            raise ValueError(f"Section {kind} must be a list of records")
        columns = set(table.c.keys())
        for row in rows:
            extra = set(row) - columns
            if extra:
                raise ValueError(f"Unknown {kind} fields: " + ", ".join(sorted(extra)))
            if kind in ENTITIES or kind == "story_image":
                if not isinstance(row.get("id"), str) or row["id"] in seen:
                    raise ValueError(f"Every {kind} needs a unique id")
                seen.add(row["id"])
        clean[kind] = [
            {k: v for k, v in row.items() if not (kind == "story" and k in ADMIN)} for row in rows
        ]
    if len(clean["story"]) != 1:
        raise ValueError("A backup holds exactly one story")
    return clean


def slug(title: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (title or "story").lower()).strip("-") or "story"


_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")


def add_runs(paragraph: Any, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif len(part) > 2 and part[0] == part[-1] and part[0] in "*_":
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def manuscript_docx(markdown: str, author_name: str | None = None, images: ImageFiles | None = None) -> bytes:
    """Lay out the compiled Markdown manuscript as a plain Word document."""
    document = Document()
    document.core_properties.author = author_name or ""
    for block in blocks(markdown):
        if block.kind == "image":
            image = (images or {}).get(block.image_id)
            if image is None:
                raise ClientException(detail="A manuscript image is unavailable")
            from PIL import Image

            with Image.open(io.BytesIO(image[0])) as picture:
                scale = min(6 / picture.width, 8 / picture.height)
                shape = document.add_picture(io.BytesIO(image[0]), width=Inches(picture.width * scale),
                                             height=Inches(picture.height * scale))
                shape._inline.docPr.set("descr", image[2] or block.text)
            if image[3]:
                document.add_paragraph(image[3])
            continue
        if block.kind == "heading":
            document.add_heading(block.text, 0 if block.level == 1 else block.level - 1)
        elif block.kind == "quote":
            quote = document.add_paragraph(style="Quote")
            add_runs(quote, block.text)
        elif block.kind == "break":
            document.add_paragraph("* * *").alignment = 1  # centred scene break
        else:
            add_runs(document.add_paragraph(), block.text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class StoryBackupController(Controller):
    path = "/api/stories/{story_id:uuid}"
    tags = ["backup"]  # noqa: RUF012 - Litestar controller configuration
    signature_namespace = {"AsyncSession": AsyncSession}  # noqa: RUF012

    @get("/backup", summary="Download the whole story as a JSON backup")
    async def backup(self, db_session: AsyncSession, story_id: UUID) -> Response[dict[str, Any]]:
        story = await db_session.get(Story, story_id)
        if story is None:
            raise NotFoundException(detail=f"No story with id {story_id}")
        exported = datetime.now(UTC)
        from storytool.domain.story.images import portable_images

        state = await full_state(db_session, story_id)
        files = await portable_images(db_session, state, include_history=True)
        return Response(
            content={
                "format": BACKUP_FORMAT,
                "version": 1,
                "exported_at": exported.isoformat(),
                "title": story.title,
                "state": state,
                "images": files,
            },
            headers={
                "content-disposition": (
                    f'attachment; filename="{slug(story.title)}-{exported:%Y%m%d}.storytool.json"'
                )
            },
        )

    @get("/manuscript.docx", summary="Compile the manuscript as a Word document")
    async def manuscript(
        self, db_session: AsyncSession, story_id: UUID, include_unplaced: bool = True,
        cover_page: bool = False,
    ) -> Response[bytes]:
        graph = await load_graph_by_id(db_session, story_id)
        if graph is None:
            raise NotFoundException(detail=f"No story with id {story_id}")
        author_name = await db_session.scalar(
            select(User.author_name).where(User.id == graph.story.user_id)
        )
        from storytool.domain.story.images import export_images

        markdown = compile_manuscript(graph, include_unplaced=include_unplaced, author_name=author_name)
        if cover_page:
            from storytool.domain.story.images import story_cover

            if cover := await story_cover(db_session, story_id):
                markdown = f"![{cover.alt}](storytool-image:{cover.id})\n\n" + markdown
        return Response(
            content=manuscript_docx(
                markdown,
                author_name,
                await export_images(db_session, story_id, markdown),
            ),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "content-disposition": f'attachment; filename="{slug(graph.story.title)}.docx"'
            },
        )

    async def _manuscript(
        self, db_session: AsyncSession, story_id: UUID, include_unplaced: bool,
        *, include_planning: bool = True,
    ) -> tuple[str, str, UUID, str | None]:
        graph = await load_graph_by_id(db_session, story_id)
        if graph is None:
            raise NotFoundException(detail=f"No story with id {story_id}")
        author_name = await db_session.scalar(
            select(User.author_name).where(User.id == graph.story.user_id)
        )
        markdown = compile_manuscript(
            graph, include_unplaced=include_unplaced, author_name=author_name,
            include_planning=include_planning,
        )
        return markdown, graph.story.title or "Untitled", graph.story.id, author_name

    @get("/manuscript.html", summary="Compile the manuscript as a standalone HTML page")
    async def export_html(
        self, db_session: AsyncSession, story_id: UUID, include_unplaced: bool = True
    ) -> Response[str]:
        markdown, title, _, author = await self._manuscript(db_session, story_id, include_unplaced)
        from storytool.domain.story.images import export_images
        return Response(
            content=manuscript_html(markdown, title, author, await export_images(db_session, story_id, markdown)),
            media_type="text/html; charset=utf-8",
            headers={"content-disposition": f'attachment; filename="{slug(title)}.html"'},
        )

    @get("/manuscript.txt", summary="Compile the manuscript as plain text")
    async def export_text(
        self, db_session: AsyncSession, story_id: UUID, include_unplaced: bool = True
    ) -> Response[str]:
        markdown, title, _, _ = await self._manuscript(db_session, story_id, include_unplaced)
        from storytool.domain.story.image_models import StoryImage
        from storytool.domain.story.images import IMAGE_REF

        descriptions = {str(image.id): image for image in await db_session.scalars(
            select(StoryImage).where(StoryImage.story_id == story_id)
        )}
        def describe_image(match: re.Match[str]) -> str:
            image = descriptions.get(match[2])
            alt = (image.alt if image else "") or match[1]
            caption = image.caption if image else ""
            return f"\n\n[Image: {alt}]\n\n" + (caption + "\n\n" if caption else "")
        markdown = IMAGE_REF.sub(describe_image, markdown)
        return Response(
            content=manuscript_text(markdown),
            media_type="text/plain; charset=utf-8",
            headers={"content-disposition": f'attachment; filename="{slug(title)}.txt"'},
        )

    @get("/manuscript.epub", summary="Compile the manuscript as an EPUB e-book")
    async def export_epub(
        self, db_session: AsyncSession, story_id: UUID, include_unplaced: bool = True
    ) -> Response[bytes]:
        markdown, title, identifier, author = await self._manuscript(db_session, story_id, include_unplaced)
        from storytool.domain.story.images import cover_file, export_images

        images = await export_images(db_session, story_id, markdown)
        cover = await cover_file(db_session, story_id)
        if cover:
            images[cover[0]] = cover[1]
        return Response(
            content=manuscript_epub(markdown, title, identifier, author, images, cover[0] if cover else None),
            media_type="application/epub+zip",
            headers={"content-disposition": f'attachment; filename="{slug(title)}.epub"'},
        )


    @get("/manuscript.pdf", summary="Submission manuscript PDF with optional reference appendices")
    async def export_pdf(
        self, db_session: AsyncSession, story_id: UUID, include_unplaced: bool = True,
        glossary: bool = False, timeline: bool = False, places: bool = False,
        cover_page: bool = False,
    ) -> Response[bytes]:
        markdown, title, _, author = await self._manuscript(
            db_session, story_id, include_unplaced, include_planning=False
        )
        graph = await load_graph_by_id(db_session, story_id)
        if graph is None:
            raise NotFoundException(detail="Story not found")
        appendices: list[tuple[str, list[Block]]] = []
        if glossary:
            entries = await db_session.scalars(
                select(GlossaryEntry).where(GlossaryEntry.story_id == story_id)
            )
            section = []
            for entry in sorted(entries, key=lambda item: (item.term.casefold(), str(item.id))):
                section.append(Block("heading", entry.term))
                if entry.aliases:
                    section.append(Block("para", "Also known as: " + ", ".join(entry.aliases)))
                if entry.definition:
                    section.append(Block("para", entry.definition))
            appendices.append(("Glossary", section or [Block("para", "No glossary entries yet.")]))
        if timeline:
            data = await story_timeline(db_session, story_id)
            section = [Block("para", "World chronology, separate from the order scenes are read.")]
            for entry in sorted(
                data.entries,
                key=lambda item: (
                    item.ordinal is None, item.ordinal or 0, item.title.casefold(), str(item.id)
                ),
            ):
                time = entry.time_label or (
                    f"Position {entry.ordinal}" if entry.ordinal is not None else "Time not assigned"
                )
                section.append(Block("heading", f"{time}: {entry.title}"))
                details = [entry.kind.capitalize()]
                if entry.is_turning_point:
                    details.append("Turning point")
                if entry.is_flashback:
                    details.append("Flashback")
                if entry.location_name:
                    details.append(entry.location_name)
                if entry.participants:
                    details.append(", ".join(person.name for person in entry.participants))
                if entry.reading_label:
                    details.append("Reading order: " + entry.reading_label)
                section.append(Block("para", " · ".join(details)))
                if entry.summary:
                    section.append(Block("para", entry.summary))
            if not data.entries:
                section.append(Block("para", "No timeline entries yet."))
            appendices.append(("Timeline", section))
        if places:
            section = []
            for place in sorted(graph.locations, key=lambda item: (item.name.casefold(), str(item.id))):
                section.append(Block("heading", place.name))
                if place.description:
                    section.append(Block("para", place.description))
                if place.atmosphere_notes:
                    section.append(Block("para", "Atmosphere: " + place.atmosphere_notes))
            appendices.append(("Places", section or [Block("para", "No places yet.")]))
        scenes = graph.scenes if include_unplaced else graph.scenes_in_reading_order()
        try:
            from storytool.domain.story.images import cover_file, export_images

            cover = await cover_file(db_session, story_id) if cover_page else None
            content = manuscript_pdf(
                markdown, title, author, sum(scene.word_count for scene in scenes), appendices,
                await export_images(db_session, story_id, markdown),
                cover[1] if cover else None,
            )
        except ValueError as error:
            raise ClientException(detail=str(error)) from None
        return Response(
            content=content, media_type="application/pdf",
            headers={"content-disposition": f'attachment; filename="{slug(title)}.pdf"'},
        )

    @get("/manuscript.zip", summary="Portable Markdown with its image files")
    async def export_markdown_images(self, db_session: AsyncSession, story_id: UUID,
                                     include_unplaced: bool = True) -> Response[bytes]:
        from storytool.domain.story.images import IMAGE_REF, export_images

        markdown, title, _, _ = await self._manuscript(db_session, story_id, include_unplaced)
        images = await export_images(db_session, story_id, markdown)
        if not images:
            return Response(content=markdown.encode(), media_type="text/markdown",
                            headers={"content-disposition": f'attachment; filename="{slug(title)}.md"'})
        def markdown_image(match: re.Match[str]) -> str:
            image = images[match[2]]
            alt = re.sub(r"[\]\n]", " ", image[2] or match[1])
            caption = image[3]
            return f"\n\n![{alt}](images/{match[2]}.{image[1].split('/')[1]})\n\n" + (
                caption + "\n\n" if caption else ""
            )
        markdown = IMAGE_REF.sub(markdown_image, markdown)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(f"{slug(title)}.md", markdown)
            for id_, image in images.items():
                archive.writestr(f"images/{id_}.{image[1].split('/')[1]}", image[0])
        return Response(content=buffer.getvalue(), media_type="application/zip",
                        headers={"content-disposition": f'attachment; filename="{slug(title)}.zip"'})


class BackupImportController(Controller):
    path = "/api/backups"
    tags = ["backup"]  # noqa: RUF012 - Litestar controller configuration
    signature_namespace = {"AsyncSession": AsyncSession}  # noqa: RUF012

    @post(
        "/import",
        status_code=201,
        summary="Restore a backup file as a new story",
        request_max_body_size=MAX_BACKUP_BYTES,
    )
    async def import_backup(
        self, db_session: AsyncSession, request: Request, data: dict[str, Any]
    ) -> StoryOut:
        user_id: UUID = request.scope["state"]["storytool_user_id"]
        sandbox = data.get("sandbox")
        import_subject = None
        if sandbox is not None:
            if not isinstance(sandbox, dict):
                raise ClientException(detail="Invalid sandbox backup metadata")
            try:
                import_subject = uuid5(user_id, str(UUID(sandbox["import_id"])))
                edits = sandbox.get("edits", [])
                if not isinstance(edits, list) or len(edits) > 10000:
                    raise ValueError("Too many sandbox history entries")
                for edit in edits:
                    UUID(edit["scene_id"])
                    datetime.fromisoformat(edit["created_at"])
            except (ValueError, TypeError, KeyError) as error:
                raise ClientException(detail="Invalid sandbox backup metadata") from error
            # Serialize retries for this user's sandbox; the marker and story commit together.
            await db_session.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:key, 0))"), {"key": str(import_subject)})
            from storytool.domain.site.funnel import FunnelEvent

            prior = await db_session.scalar(select(Story).join(FunnelEvent, FunnelEvent.id == Story.id).where(FunnelEvent.kind == "try_restore", FunnelEvent.subject_id == import_subject, Story.user_id == user_id))
            if prior is not None:
                if prior.deleted_at is not None:
                    raise ClientException(detail="Your kept sandbox story is in Trash. Restore it from your library.")
                return StoryOut.model_validate(prior)
        try:
            state = validated_state(data)
            files = data.get("images", {})
            if not isinstance(files, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in files.items()):
                raise ValueError("Invalid image files in backup")
            from storytool.domain.story.images import import_images

            image_ids = {str(UUID(k)): str(uuid7()) for k in files}
            if set(image_ids) & {r["id"] for kind in ENTITIES for r in state[kind]}:
                raise ValueError("Image IDs overlap story entity IDs")
            async with db_session.begin_nested():
                created = Story(
                    title=state["story"][0].get("title") or "Restored story", user_id=user_id
                )
                db_session.add(created)
                await db_session.flush()
                copied = rehome(state, created.id, strict=True, image_ids=image_ids)
                await import_images(db_session, state, files, copied, created.id, image_ids=image_ids)
                await restore_state(db_session, created.id, copied)
        except (ValueError, TypeError, KeyError, SQLAlchemyError) as error:
            detail = str(error) if isinstance(error, ValueError) else "The backup is damaged"
            raise ClientException(detail=f"Could not restore backup: {detail}") from error
        await db_session.refresh(created)
        from storytool.domain.story.authorship import record_changes

        await record_changes(
            db_session,
            created.id,
            user_id,
            {},
            await full_state(db_session, created.id),
            "import",
            "Restored from backup",
        )
        if sandbox is not None:
            from storytool.domain.site.funnel import FunnelEvent, record_milestone, request_visitor
            from storytool.domain.story.authorship import FieldAuthorship

            scene_ids = {old["id"]: new["id"] for old, new in zip(state["scene"], copied["scene"], strict=True)}
            for edit in sandbox.get("edits", []):
                new_id = scene_ids.get(edit["scene_id"])
                if new_id is None:
                    continue  # An edited scene may have been deleted in free play.
                moment = datetime.fromisoformat(edit["created_at"])
                if moment.tzinfo is None:
                    moment = moment.replace(tzinfo=UTC)
                moment = min(moment, datetime.now(UTC))
                db_session.add(FieldAuthorship(story_id=created.id, user_id=user_id, entity_type="scene", entity_id=UUID(new_id), field="content", origin="author", actor_label="You", action="update", created_at=moment))
            db_session.add(FunnelEvent(id=created.id, kind="try_restore", subject_id=import_subject, user_id=user_id, visitor_id=request_visitor(request)))
            await record_milestone(db_session, "try_imported", user_id=user_id, visitor_id=request_visitor(request))
        await checkpoint(db_session, created.id, user_id, "Kept from sandbox" if sandbox else "Restored from backup", "initial")
        return StoryOut.model_validate(created)
